"""Excel import, validation (S1) and export (S11).

Every cell is read as text, which matches ``pandas.read_excel(..., dtype=str)``:
IDs and pincodes keep leading zeros and are never coerced to numbers.
The export always writes a *copy* of the workbook; the source is never touched.
"""

import datetime as _dt
import math
import os
import re
import shutil
import tempfile

from . import constants as C
from .config import get

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError:  # pragma: no cover
    openpyxl = None


class ExcelSchemaError(ValueError):
    """The workbook does not have the expected sheet or columns."""


def _require_openpyxl():
    if openpyxl is None:
        raise ExcelSchemaError("openpyxl is not available in this Python environment")


# ---------------------------------------------------------------------------
# Cell conversion
# ---------------------------------------------------------------------------

def cell_to_text(value):
    """Convert a cell value to text exactly as pandas dtype=str would present it."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return str(value)
    return str(value).strip()


_NUM_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?$")


def parse_coordinate(text):
    """Parse a decimal-degree coordinate from text. Returns float or None."""
    if text is None:
        return None
    t = str(text).strip().replace("°", "").replace(" ", "")
    t = t.rstrip("NnEeSsWw")
    if t.count(",") == 1 and "." not in t:
        t = t.replace(",", ".")          # decimal comma
    if not _NUM_RE.match(t):
        return None
    v = float(t)
    if math.isnan(v) or math.isinf(v):
        return None
    return v


def _loose(text):
    """Header key ignoring case, spacing and non-breaking spaces ("Survey  id " == "Survey ID")."""
    t = str(text).replace(" ", " ")
    for q in ("’", "‘", "`", "�"):     # curly / broken apostrophes
        t = t.replace(q, "'")
    return " ".join(t.lower().split())


def _col_ref_index(ref):
    """'@AB' -> 28 (1-based)."""
    return column_index_from_string(ref[1:].strip().upper())


# ---------------------------------------------------------------------------
# Header mapping
# ---------------------------------------------------------------------------

class HeaderMap(object):
    """Maps header text and column letters to 1-based column indexes."""

    def __init__(self, headers):
        # headers: list of (col_index, text)
        self.headers = headers
        self.first = {}
        self.loose = {}  # case/space-insensitive fallback
        self.keys = {}  # col_index -> unique key used in source JSON
        seen = {}
        for idx, text in headers:
            norm = (text or "").strip()
            if norm and norm not in self.first:
                self.first[norm] = idx
            if norm and _loose(norm) not in self.loose:
                self.loose[_loose(norm)] = idx
            base = norm or "Column %s" % get_column_letter(idx)
            seen[base] = seen.get(base, 0) + 1
            self.keys[idx] = base if seen[base] == 1 else "%s (%d)" % (base, seen[base])

    def index_of(self, ref):
        if not ref:
            return None
        if ref.startswith("@"):
            return _col_ref_index(ref)
        exact = self.first.get(ref.strip())
        if exact is not None:
            return exact
        hit = self.loose.get(_loose(ref))
        if hit is not None:
            return hit
        for idx, key in self.keys.items():    # unique keys, e.g. "Remark (2)" for a repeated header
            if _loose(key) == _loose(ref):
                return idx
        return None

    def field_index(self, field, header, aliases=None):
        """Column of an internal field: the configured header first, then its aliases."""
        idx = self.index_of(header)
        if idx is None:
            for alt in (aliases or {}).get(field, []) or []:
                idx = self.index_of(alt)
                if idx is not None:
                    break
        return idx

    @property
    def max_col(self):
        return max([i for i, _ in self.headers] or [0])


REQUIRED_FIELDS = ("survey_id", "latitude", "longitude")


def _pick_sheet(wb, cfg, header_row):
    """The configured sheet, or else the only/first sheet that has the required columns.

    Returns (worksheet, header_map, note). A client file whose sheet is called
    "Data" instead of "Sheet1" should still import - the note says what happened.
    """
    wanted = get(cfg, "excel.sheet_name", "Sheet1")
    columns = get(cfg, "excel.columns", {}) or {}
    aliases = get(cfg, "excel.column_aliases", {}) or {}
    if wanted in wb.sheetnames:
        ws = wb[wanted]
        return ws, _read_header(ws, header_row), None
    for name in wb.sheetnames:
        ws = wb[name]
        hmap = _read_header(ws, header_row)
        if all(hmap.field_index(f, columns.get(f), aliases) is not None for f in REQUIRED_FIELDS):
            return ws, hmap, "Sheet %r not found; used sheet %r which has the survey columns" % (wanted, name)
    raise ExcelSchemaError("No sheet named %r and no sheet with Survey ID / latitude / longitude columns "
                           "(sheets: %s)" % (wanted, ", ".join(wb.sheetnames)))


def _read_header(ws, header_row):
    headers = []
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        for idx, value in enumerate(row, start=1):
            headers.append((idx, cell_to_text(value) or None))
    # Trim trailing empty header columns
    while headers and headers[-1][1] is None:
        headers.pop()
    return HeaderMap(headers)


# ---------------------------------------------------------------------------
# Import + validation
# ---------------------------------------------------------------------------

def read_records(path, cfg):
    """Read the workbook and return ``(records, report)``.

    Each record is a dict with survey_id, row, latitude, longitude, fields,
    source, state (IMPORTED or INVALID), reasons (invalid) and warnings.
    """
    _require_openpyxl()
    if not os.path.isfile(path):
        raise ExcelSchemaError("Workbook not found: %s" % path)
    header_row = int(get(cfg, "excel.header_row", 1))
    columns = get(cfg, "excel.columns", {}) or {}
    aliases = get(cfg, "excel.column_aliases", {}) or {}
    control_refs = get(cfg, "excel.control_columns", []) or []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws, hmap, sheet_note = _pick_sheet(wb, cfg, header_row)
        sheet_name = ws.title

        field_cols = {}
        missing = []
        for field, header in columns.items():
            idx = hmap.field_index(field, header, aliases)
            if idx is None:
                missing.append((field, header))
            else:
                field_cols[field] = idx
        required_missing = ["%s (%s)" % (h, f) for f, h in missing
                            if f in ("survey_id", "latitude", "longitude")]
        if required_missing:
            raise ExcelSchemaError("Required columns missing: %s" % ", ".join(required_missing))

        control_idx = set()
        for ref in control_refs:
            i = hmap.index_of(ref)
            if i:
                control_idx.add(i)
        data_cols = [(i, hmap.keys[i]) for i, _ in hmap.headers if i not in control_idx]

        records = []
        for excel_row, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                        start=header_row + 1):
            values = {}
            for i, key in data_cols:
                v = row[i - 1] if i - 1 < len(row) else None
                values[i] = cell_to_text(v)
            if not any(values.values()):
                continue  # fully empty data row (control columns ignored)
            fields = {f: values.get(i, "") for f, i in field_cols.items()}
            source = {key: values.get(i, "") for i, key in data_cols}
            records.append({
                "survey_id": fields.get("survey_id", ""),
                "row": excel_row,
                "latitude": None,
                "longitude": None,
                "fields": fields,
                "source": source,
                "state": C.IMPORTED,
                "reasons": [],
                "warnings": [],
            })
    finally:
        wb.close()

    _validate(records, cfg)
    report = {
        "path": os.path.abspath(path),
        "sheet": sheet_name,
        "total": len(records),
        "imported": sum(1 for r in records if r["state"] == C.IMPORTED),
        "invalid": sum(1 for r in records if r["state"] == C.INVALID),
        "warnings": sum(1 for r in records if r["warnings"]),
        "missing_optional_columns": ["%s (%s)" % (h, f) for f, h in missing],
        "headers": [h for _, h in hmap.headers],
        "notes": [n for n in [sheet_note] if n],
    }
    return records, report


def _validate(records, cfg):
    b = get(cfg, "project.country_bounds", {}) or {}
    min_lat, max_lat = float(b.get("min_lat", -90)), float(b.get("max_lat", 90))
    min_lon, max_lon = float(b.get("min_lon", -180)), float(b.get("max_lon", 180))

    def inside(lat, lon):
        return min_lat <= lat <= max_lat and min_lon <= lon <= max_lon

    id_rows = {}
    for r in records:
        sid = r["survey_id"]
        if sid:
            id_rows.setdefault(sid, []).append(r["row"])

    loc_rows = {}
    for r in records:
        f = r["fields"]
        reasons, warnings = r["reasons"], r["warnings"]
        sid = r["survey_id"]
        if not sid:
            reasons.append("Survey ID is empty")
        elif len(id_rows[sid]) > 1:
            reasons.append("Duplicate Survey ID (rows %s)" % ", ".join(str(x) for x in id_rows[sid]))

        lat = parse_coordinate(f.get("latitude"))
        lon = parse_coordinate(f.get("longitude"))
        if lat is None:
            reasons.append("Latitude is missing or not numeric (%r)" % f.get("latitude", ""))
        if lon is None:
            reasons.append("Longitude is missing or not numeric (%r)" % f.get("longitude", ""))
        if lat is not None and lon is not None:
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                if -90 <= lon <= 90 and inside(lon, lat):
                    reasons.append("Latitude/longitude appear swapped")
                else:
                    reasons.append("Coordinates are not valid degrees (%s, %s)" % (lat, lon))
            elif not inside(lat, lon):
                if inside(lon, lat):
                    reasons.append("Latitude/longitude appear swapped")
                else:
                    reasons.append("Coordinates outside the configured country bounds (%s, %s)" % (lat, lon))
            else:
                r["latitude"], r["longitude"] = lat, lon
                key = (round(lat, 6), round(lon, 6))
                loc_rows.setdefault(key, []).append(r["row"])
                decimals = min(_decimals(f.get("latitude")), _decimals(f.get("longitude")))
                if decimals < 4:
                    warnings.append("Low coordinate precision (%d decimals, ~%d m)" % (decimals, 11 * 10 ** (4 - decimals)))

        if not f.get("building_name"):
            warnings.append("Building name is empty")
        addr_parts = [f.get(k, "") for k in ("street", "area", "city", "pincode", "standard_address")]
        if not any(addr_parts):
            warnings.append("Address is empty")
        if reasons:
            r["state"] = C.INVALID

    for rows in loc_rows.values():
        if len(rows) > 1:
            for r in records:
                if r["row"] in rows and r["state"] == C.IMPORTED:
                    others = [str(x) for x in rows if x != r["row"]]
                    r["warnings"].append("Same coordinates as row(s) %s" % ", ".join(others))


def _decimals(text):
    t = str(text or "").strip()
    if "." not in t:
        return 0
    return len(t.split(".", 1)[1].rstrip())


# ---------------------------------------------------------------------------
# Remark / client columns
# ---------------------------------------------------------------------------

def choose_remark(cfg, shift_m, dist_to_polygon_m):
    """Pick the client's remark text from the corrected-coordinate evidence.

    shift_m: distance original -> corrected coordinate (metres)
    dist_to_polygon_m: distance original point -> final polygon (0 if inside)
    """
    remarks = get(cfg, "excel.remarks", {}) or {}
    rules = get(cfg, "excel.remark_rules", {}) or {}
    if shift_m is not None and shift_m <= float(rules.get("unchanged_max_shift_m", 1.0)):
        return remarks.get("unchanged", "")
    if dist_to_polygon_m is not None and dist_to_polygon_m <= float(rules.get("near_max_distance_m", 30.0)):
        return remarks.get("adjusted_near", "")
    return remarks.get("modified", "")


def format_latlon(cfg, lat, lon):
    fmt = get(cfg, "excel.corrected_latlon_format", "{lat:.6f},{lon:.6f}")
    return fmt.format(lat=lat, lon=lon)


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def export_workbook(source_path, output_path, results, cfg, exported_at=None,
                    exportable_statuses=C.EXPORTABLE_STATUSES):
    """Write a copy of the source workbook with results filled in.

    ``results`` maps survey_id -> dict with keys: status, corrected_lat,
    corrected_lon, polygon_id, area_m2, method, confidence, route, reviewer,
    reviewed_at, strategy, shift_m, dist_to_polygon_m, footprint_source,
    imagery_source, config_version, model_version, manifest_hash.

    Only results whose status is in ``exportable_statuses`` and that carry a
    corrected coordinate are written. Returns a report dict.
    """
    _require_openpyxl()
    src = os.path.abspath(source_path)
    out = os.path.abspath(output_path)
    if os.path.normcase(src) == os.path.normcase(out):
        raise ExcelSchemaError("Refusing to overwrite the source workbook")
    exported_at = exported_at or _dt.datetime.now().replace(microsecond=0)

    header_row = int(get(cfg, "excel.header_row", 1))
    wb = openpyxl.load_workbook(src)  # full mode keeps styles, merges, widths
    ws, hmap, _ = _pick_sheet(wb, cfg, header_row)

    sid_col = hmap.field_index("survey_id", get(cfg, "excel.columns.survey_id"),
                               get(cfg, "excel.column_aliases", {}) or {})
    if sid_col is None:
        raise ExcelSchemaError("Survey ID column not found")

    client_cols = {}
    for key, header in (get(cfg, "excel.client_output_columns", {}) or {}).items():
        idx = hmap.index_of(header)
        if idx is not None:
            client_cols[key] = idx

    # Appended output columns go after the last used template column.
    appended = get(cfg, "excel.appended_columns", {}) or {}
    last_col = max(hmap.max_col, ws.max_column)
    header_src = ws.cell(row=header_row, column=hmap.max_col)
    app_cols = {}
    for n, (key, title) in enumerate(appended.items(), start=1):
        col = last_col + n
        app_cols[key] = col
        c = ws.cell(row=header_row, column=col, value=title)
        if header_src.has_style:
            c._style = header_src._style
        ws.column_dimensions[get_column_letter(col)].width = max(14, min(40, len(title) + 4))

    date_fmt = get(cfg, "excel.delivery_date_format", "DD-MM-YYYY")
    status_values = get(cfg, "excel.status_values", {}) or {}
    decimals = int(get(cfg, "corrected_coordinate.decimals", 6))

    written, skipped, seen = [], [], set()
    for r in range(header_row + 1, ws.max_row + 1):
        sid = cell_to_text(ws.cell(row=r, column=sid_col).value)
        if not sid:
            continue
        seen.add(sid)
        res = results.get(sid)
        if not res:
            continue
        status = res.get("status")
        lat, lon = res.get("corrected_lat"), res.get("corrected_lon")
        if status not in exportable_statuses or lat is None or lon is None:
            skipped.append((sid, status or "NO_RESULT"))
            continue
        lat, lon = round(float(lat), decimals), round(float(lon), decimals)

        def put(col_map, key, value, number_format=None):
            col = col_map.get(key)
            if col is None:
                return
            cell = ws.cell(row=r, column=col, value=value)
            if number_format:
                cell.number_format = number_format

        put(client_cols, "remark", choose_remark(cfg, res.get("shift_m"), res.get("dist_to_polygon_m")))
        put(client_cols, "corrected_latlon", format_latlon(cfg, lat, lon))
        put(client_cols, "operator_name", res.get("reviewer") or "")
        put(client_cols, "status", status_values.get(status, status))
        put(client_cols, "delivery_date", exported_at, date_fmt)

        conf = res.get("confidence")
        area = res.get("area_m2")
        put(app_cols, "corrected_latitude", lat, "0.000000")
        put(app_cols, "corrected_longitude", lon, "0.000000")
        put(app_cols, "polygon_id", res.get("polygon_id") or "")
        put(app_cols, "area_m2", round(float(area), 1) if area is not None else None, "0.0")
        put(app_cols, "method", res.get("method") or "")
        put(app_cols, "confidence", round(float(conf), 3) if conf is not None else None, "0.000")
        put(app_cols, "route", res.get("route") or "")
        put(app_cols, "review_status", status)
        put(app_cols, "reviewer", res.get("reviewer") or "")
        put(app_cols, "reviewed_at", _as_datetime(res.get("reviewed_at")), "YYYY-MM-DD HH:MM:SS")
        put(app_cols, "exported_at", exported_at, "YYYY-MM-DD HH:MM:SS")
        put(app_cols, "coordinate_strategy", res.get("strategy") or "")
        put(app_cols, "footprint_source", res.get("footprint_source") or "")
        put(app_cols, "imagery_source", res.get("imagery_source") or "")
        put(app_cols, "config_version", res.get("config_version") or "")
        put(app_cols, "model_version", res.get("model_version") or "")
        put(app_cols, "manifest_hash", res.get("manifest_hash") or "")
        written.append(sid)

    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(out) or ".")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
    return {
        "output_path": out,
        "exported": len(written),
        "exported_ids": written,
        "skipped": skipped,
        "unknown_ids": sorted(set(results) - seen),
        "appended_columns": {k: get_column_letter(v) for k, v in app_cols.items()},
        "exported_at": exported_at.isoformat(),
    }


def _as_datetime(value):
    if value is None or value == "":
        return None
    if isinstance(value, _dt.datetime):
        return value
    try:
        return _dt.datetime.fromisoformat(str(value).replace("Z", ""))
    except ValueError:
        return str(value)


ENTRY_FIELDS = ("survey_id", "latitude", "longitude", "customer_name", "building_name", "building_no", "tower_block",
                "street", "area", "landmark", "city", "district", "state", "pincode", "building_category", "floors",
                "building_height_m", "building_type")


def write_entry_workbook(cfg, entries, path):
    """Write buildings entered one by one as a workbook in the configured template layout.

    The result is imported like any client workbook, so validation, processing, review,
    export and the audit trail are exactly the same as for an Excel batch. ``entries``:
    dicts keyed by field name (survey_id, latitude, longitude, building_name, ...).
    """
    _require_openpyxl()
    if not entries:
        raise ExcelSchemaError("No buildings to write")
    columns = get(cfg, "excel.columns", {}) or {}
    headers = [(f, h) for f, h in columns.items() if isinstance(h, str) and not h.startswith("@")]
    for f in ("survey_id", "latitude", "longitude"):
        if f not in dict(headers):
            raise ExcelSchemaError("The config has no header for %s" % f)
    extra = [h for h in (get(cfg, "excel.client_output_columns", {}) or {}).values()
             if h and h not in dict(headers).values()]
    header_row = int(get(cfg, "excel.header_row", 1))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = str(get(cfg, "excel.sheet_name", "Sheet1") or "Sheet1")[:31]
    titles = [h for _, h in headers] + extra
    for col, title in enumerate(titles, start=1):
        ws.cell(row=header_row, column=col, value=title)
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(40, len(title) + 2))
    for i, e in enumerate(entries, start=header_row + 1):
        for col, (field, _h) in enumerate(headers, start=1):
            v = e.get(field)
            if v is None or v == "":
                continue
            if field in ("latitude", "longitude"):
                ws.cell(row=i, column=col, value=float(v)).number_format = "0.0000000"
            else:
                ws.cell(row=i, column=col, value=str(v))
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def copy_input(source_path, batch_dir):
    """Copy the workbook into batch_<id>/input/ and mark it read-only."""
    dest_dir = os.path.join(batch_dir, "input")
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, os.path.basename(source_path))
    if not os.path.exists(dest):
        shutil.copy2(source_path, dest)
        try:
            os.chmod(dest, 0o444)
        except OSError:
            pass
    return dest
