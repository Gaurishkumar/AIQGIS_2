"""Plugin test of the imagery + SAM path and click-to-fix, on a synthetic licensed ortho.

  worker\\.venv\\Scripts\\python.exe tests\\qgis\\make_synthetic_project.py <proj>
  "C:\\Program Files\\QGIS 4.2.2\\bin\\python-qgis.bat" tests\\qgis\\test_plugin_imagery.py <proj>
"""

import json
import os
import sys
import tempfile

from harness import ROOT, Check, FakeIface, app, pump, wait_until

app()

from qgis.core import QgsGeometry  # noqa: E402

from building_polygon_automation import geometry_ops as G  # noqa: E402
from building_polygon_automation.bpa_common import constants as C  # noqa: E402
from building_polygon_automation.compat import geom_equal, val  # noqa: E402
from building_polygon_automation.plugin import BuildingPolygonAutomationPlugin  # noqa: E402


def iou(a, b, epsg):
    a, b = G.to_utm(a, epsg), G.to_utm(b, epsg)
    inter = a.intersection(b).area()
    union = a.combine(b).area()
    return inter / union if union else 0.0


def main(proj):
    check = Check()
    truth = {f["properties"]["survey_id"]: G.geojson_to_geometry(f["geometry"])
             for f in json.load(open(os.path.join(proj, "truth.geojson")))["features"]}
    iface = FakeIface()
    plugin = BuildingPolygonAutomationPlugin(iface)
    plugin.initGui()
    s = plugin.settings
    s.set("project_dir", proj)
    s.set("worker_python", os.path.join(ROOT, "worker", ".venv", "Scripts", "python.exe"))
    s.set("batches_dir", tempfile.mkdtemp(prefix="bpa_syn_batches_"))
    s.set("operator_id", "tester02")
    s.set("port", 8793)
    s.set("auto_next", False)
    s.set("load_model", True)
    plugin.toggle(True)
    c = plugin.controller
    check(s.worker_dir().endswith("worker"), "worker folder derived from the worker Python")

    rep = c.import_excel(os.path.join(proj, "synthetic_survey.xlsx"))
    check(rep["total"] == 6, "6 synthetic records imported")
    c.process()
    wait_until(lambda: c.worker_ready, 240, "worker health")
    imagery = c.health.get("imagery_sources", [])
    check(imagery and imagery[0]["enabled"], "licensed imagery source enabled in /health")
    wait_until(lambda: c.worker_batch_id is not None, 60, "submission")
    wait_until(lambda: not c.poll_timer.isActive() and len(c.details) >= 6, 400, "SAM batch")
    wait_until(lambda: c.health.get("model", {}).get("loaded") is True, 30, "health refresh")
    model = c.health.get("model", {})
    check(model.get("loaded") is True, "SAM 2.1 loaded in the worker (%s)" % model.get("model_version"))
    mv = [val(f["model_version"]) for f in c.store.all_buildings()]
    check(all(v and v.startswith("sam2.1_hiera") for v in mv), "model version stored on every SAM building")
    check(all(val(f["manifest_hash"]) for f in c.store.all_buildings()), "manifest hash stored on every building")

    ious_fp, ious_prop = [], []
    for key, r in c.records.items():
        d = c.details[key]
        p = d["proposal"]
        epsg = d["utm_epsg"]
        chosen = next(x for x in d["candidates"] if x["chosen"])
        fp = G.geojson_to_geometry(chosen["geometry"])
        prop = c.store.building(r["survey_id"]).geometry()
        ious_fp.append(iou(fp, truth[r["survey_id"]], epsg))
        ious_prop.append(iou(prop, truth[r["survey_id"]], epsg))
        print("   %s %-14s IoU footprint %.3f -> proposal %.3f" % (r["survey_id"], p["method"], ious_fp[-1], ious_prop[-1]))
    check(all("SAM" in c.details[k]["proposal"]["method"] for k in c.records), "every record refined with SAM")
    med = sorted(ious_prop)[len(ious_prop) // 2]
    check(med >= 0.85, "median proposal IoU vs truth >= 0.85 (%.3f)" % med)
    check(sum(ious_prop) > sum(ious_fp), "SAM proposals beat raw footprints (%.3f vs %.3f mean)" % (
        sum(ious_prop) / 6, sum(ious_fp) / 6))

    key = "SYN_B04"
    c.select(key)
    check(c.crop_layer is not None and c.crop_layer.isValid(), "imagery crop layer loaded for the record")
    d = c.details[key]
    epsg = d["utm_epsg"]

    print("\n== click-to-fix")
    t = truth[key]
    tank = t.pointOnSurface().asPoint()     # the synthetic roof has a water tank exactly here
    ctr = t.centroid().asPoint()
    far = G.to_wgs84(G.to_utm(t, epsg).buffer(6, 4), epsg).boundingBox()
    rev = d["revision"]
    c.click_to_fix([[tank.x(), tank.y()]], [[far.xMinimum(), far.yMinimum()]], None, key)
    wait_until(lambda: c.details[key]["revision"] != rev, 90, "resegment")
    p = c.details[key]["proposal"]
    b = c.store.building("SYN_B04")
    check(p["method"] == C.METHOD_CLICK_TO_FIX and val(b["method"]) == C.METHOD_CLICK_TO_FIX, "click-to-fix replaced the proposal")
    check(val(b["status"]) == C.STATUS_PROPOSED, "click-to-fix result waits for approval")
    fix_iou = iou(b.geometry(), t, epsg)
    check(fix_iou > 0.85, "single click on a roof object still gives the whole roof (IoU %.3f)" % fix_iou)
    check(any(a["action"] == C.ACTION_CLICK_TO_FIX for a in c.store.actions("SYN_B04")), "click-to-fix audited")

    print("\n== refused click-to-fix keeps the polygon")
    before = QgsGeometry(c.store.building("SYN_B04").geometry())
    tiny = G.to_wgs84(G.to_utm(QgsGeometry.fromPointXY(ctr), epsg).buffer(0.8, 4), epsg).boundingBox()
    rev = c.details[key]["revision"]
    n_warn = len(iface.bar.messages)
    c.click_to_fix([], [], [tiny.xMinimum(), tiny.yMinimum(), tiny.xMaximum(), tiny.yMaximum()], key)
    wait_until(lambda: c.details[key]["revision"] != rev, 90, "tiny box resegment")
    check(c.details[key]["proposal"]["route"] == C.ROUTE_MANUAL, "tiny mask fails the area gate")
    check(geom_equal(c.store.building("SYN_B04").geometry(), before), "previous polygon kept after a refused fix")
    check(any("refused" in m[1] for m in iface.bar.messages[n_warn:]), "operator told why the fix was refused")
    bb = G.to_wgs84(G.to_utm(t, epsg).buffer(1.5, 4), epsg).boundingBox()
    rev = c.details[key]["revision"]
    c.click_to_fix([[ctr.x(), ctr.y()]], [], [bb.xMinimum(), bb.yMinimum(), bb.xMaximum(), bb.yMaximum()], key)
    wait_until(lambda: c.details[key]["revision"] != rev, 90, "box resegment")
    box_iou = iou(c.store.building("SYN_B04").geometry(), t, epsg)
    check(box_iou > 0.85, "box prompt IoU vs truth %.3f" % box_iou)
    seg = c.details[key]["segmentation"]["click_to_fix"]
    check(seg["box"] is True, "box prompt reached the worker")

    print("\n== approve all and export")
    for k in list(c.records):
        c.select(k, zoom=False)
        c.approve(k, confirm_overlap=True)
    out = os.path.join(c.batch["batch_dir"], "output", "syn_export.xlsx")
    rep = c.export(out)
    check(rep["exported"] == 6, "6 records exported")
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    hdr = {ws.cell(1, col).value: col for col in range(1, ws.max_column + 1)}
    imgsrc = [ws.cell(r, hdr["Imagery Source"]).value for r in range(2, 8)]
    check(all(v and "SYNTHETIC-TEST-001" in v for v in imgsrc), "imagery provenance (licence ref) exported")
    c.stop_worker()
    pump(1.0)
    plugin.unload()
    return check.done()


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
