"""End-to-end plugin test inside QGIS (offscreen) with a real worker process.

Run with QGIS's Python:
  "C:\\Program Files\\QGIS 4.2.2\\bin\\python-qgis.bat" tests\\qgis\\test_plugin_e2e.py <workbook.xlsx> [batches_dir]

Covers: import + validation, worker start via QProcess, coverage check,
batch processing, proposals as real GeoPackage features, approve (corrected
coordinate), reject, native edit capture (EDITED + audit), manual drawing
(APPROVED_MANUAL), pick-other-candidate, click-to-fix refusal without
imagery, export (copy of the workbook, structurally identical), audit trail.
"""

import os
import sys
import tempfile

from harness import ROOT, Check, FakeIface, app, pump, wait_until

app()

import openpyxl  # noqa: E402
from qgis.core import QgsFeature, QgsGeometry, QgsPointXY  # noqa: E402

from building_polygon_automation.bpa_common import constants as C  # noqa: E402
from building_polygon_automation.compat import val  # noqa: E402
from building_polygon_automation.controller import NeedsConfirmation, UserError  # noqa: E402
from building_polygon_automation.plugin import BuildingPolygonAutomationPlugin  # noqa: E402


def main(workbook, batches_dir=None):
    check = Check()
    batches_dir = batches_dir or tempfile.mkdtemp(prefix="bpa_batches_")
    iface = FakeIface()
    plugin = BuildingPolygonAutomationPlugin(iface)
    plugin.initGui()
    s = plugin.settings
    s.set("project_dir", ROOT)
    s.set("batches_dir", batches_dir)
    s.set("operator_id", "tester01")
    s.set("port", 8791)
    s.set("auto_next", False)
    s.set("load_model", False)
    plugin.toggle(True)
    c = plugin.controller
    c.auto_prepare_reference = False

    print("\n== import")
    rep = c.import_excel(workbook)
    check(rep["total"] == 100 and rep["invalid"] == 0, "100 records imported, none invalid")
    check(len(list(c.layers["survey_points"].getFeatures())) == 100, "survey_points layer has 100 points")
    check(c.store.meta().get("batch_id") == c.batch["batch_id"], "batch metadata written")
    check(os.path.isfile(os.path.join(c.batch["batch_dir"], "input", os.path.basename(workbook))), "read-only input copy")
    check(plugin.dock.table.rowCount() == 100, "all 100 imported rows listed straight after import")
    check(plugin.dock.filter.currentText() == "All records (100)" and "100 of 100" in plugin.dock.shown_label.text(),
          "filter shows record counts")
    warned = [r for r in c.records.values() if r["warnings"]]
    check(len(warned) == 3, "3 records carry validation warnings (visible, not dropped)")

    print("\n== process (worker via QProcess)")
    c.process()
    wait_until(lambda: c.worker_ready, 240, "worker health")
    check(c.worker_ready, "worker started and healthy")
    check(c.health.get("reference", {}).get("manifest_ok") is True, "reference manifest verified by worker")
    wait_until(lambda: c.worker_batch_id is not None, 60, "batch submission")
    wait_until(lambda: not c.poll_timer.isActive() and len(c.details) >= 100, 300, "batch results")
    pump(0.5)
    states = c.counts()
    print("   states:", states)
    check(sum(states.get(k, 0) for k in (C.REVIEW, C.MANUAL)) == 100, "all 100 records reached REVIEW/MANUAL")
    blds = list(c.layers["buildings"].getFeatures())
    check(len(blds) == states.get(C.REVIEW, 0), "one PROPOSED building feature per REVIEW record")
    check(all(val(f["status"]) == C.STATUS_PROPOSED for f in blds), "proposals stored as PROPOSED features")
    check(len(list(c.store.layer("candidates").getFeatures())) > 500, "candidates stored in review.gpkg")
    acts = c.store.actions()
    check(sum(1 for a in acts if a["action"] == C.ACTION_LOAD_PROPOSAL) == len(blds), "proposal loads audited")

    queue = c.queue("open")
    check(len(queue) == 100, "review queue lists all open records")
    confs = [c.records[k]["confidence"] for k in queue if c.records[k].get("route") == C.ROUTE_REVIEW]
    check(confs == sorted(confs), "queue sorted by confidence (lowest first) within route")

    print("\n== approve")
    k1 = queue[0]
    c.select(k1)
    check("Candidates" in plugin.dock.evidence.toPlainText() or c.evidence_html(k1), "evidence panel rendered")
    cc = c.approve(k1)
    b1 = c.store.building(c.records[k1]["survey_id"])
    check(val(b1["status"]) == C.STATUS_APPROVED, "approve -> APPROVED")
    check(val(b1["reviewer"]) == "tester01", "reviewer recorded")
    d1 = c.details[k1]["proposal"]
    check(abs(cc["corrected_lat"] - d1["corrected_lat"]) < 2e-6 and abs(cc["corrected_lon"] - d1["corrected_lon"]) < 2e-6,
          "plugin corrected coordinate matches worker (%.6f,%.6f vs %.6f,%.6f)" % (
              cc["corrected_lat"], cc["corrected_lon"], d1["corrected_lat"], d1["corrected_lon"]))
    check(c.store.corrected(c.records[k1]["survey_id"]) is not None, "corrected point stored")

    print("\n== reject")
    k2 = queue[1]
    c.select(k2)
    c.reject("WRONG_LOCATION", "test reject")
    check(c.records[k2]["state"] == C.REJECTED, "reject -> REJECTED")
    try:
        c.approve(k2)
        check(False, "rejected record cannot be approved directly")
    except UserError:
        check(True, "rejected record cannot be approved directly")

    print("\n== native edit is captured")
    k3 = queue[2]
    c.select(k3)
    sid3 = c.records[k3]["survey_id"]
    c.start_edit(k3)
    lyr = c.layers["buildings"]
    f3 = c.store.building(sid3)
    g3 = QgsGeometry(f3.geometry())
    moved = QgsGeometry(g3)
    moved.translate(0.00001, 0.0)
    check(lyr.changeGeometry(f3.id(), moved), "geometry changed in edit buffer")
    check(lyr.commitChanges(), "edits committed")
    pump(0.3)
    f3b = c.store.building(sid3)
    check(val(f3b["status"]) == C.STATUS_EDITED, "saved edit sets EDITED")
    edits = [a for a in c.store.actions(sid3) if a["action"] == C.ACTION_EDIT]
    check(len(edits) == 1 and edits[0]["geometry_before"] and edits[0]["geometry_after"], "edit audited with before/after")
    c.approve(k3)
    check(val(c.store.building(sid3)["status"]) == C.STATUS_APPROVED, "edited polygon approved")

    print("\n== overlap guard")
    dup = [k for k in queue if c.records[k]["survey_id"] in ("SVY_0084853", "SVY_0084933")]
    if len(dup) == 2:
        c.select(dup[0])
        c.approve(dup[0])
        c.select(dup[1])
        try:
            c.approve(dup[1])
            check(False, "duplicate survey of the same building asks for confirmation")
        except NeedsConfirmation as exc:
            check("same building" in str(exc), "duplicate survey of the same building asks for confirmation")
            c.approve(dup[1], confirm_overlap=True)
            check(c.records[dup[1]]["state"] == C.APPROVED, "confirmed overlap approval")

    print("\n== draw manually")
    manual = [k for k in queue if c.records[k]["state"] in (C.MANUAL, C.REVIEW)
              and c.records[k]["survey_id"] not in ("SVY_0084853", "SVY_0084933")]
    k4 = manual[-1]
    c.select(k4)
    r4 = c.records[k4]
    had_proposal = c.store.building(r4["survey_id"]) is not None
    c.start_draw(k4)
    if had_proposal:
        ids_in_edit = [f.id() for f in lyr.getFeatures() if val(f["survey_id"]) == r4["survey_id"]]
        check(ids_in_edit == [], "old proposal set aside while drawing (no clipping against it)")
    lon, lat = r4["longitude"], r4["latitude"]
    d = 0.00008
    ring = [QgsPointXY(lon - d, lat - d), QgsPointXY(lon + d, lat - d), QgsPointXY(lon + d, lat + d),
            QgsPointXY(lon - d, lat + d), QgsPointXY(lon - d, lat - d)]
    nf = QgsFeature(lyr.fields())
    nf.setGeometry(QgsGeometry.fromPolygonXY([ring]))
    check(lyr.addFeature(nf), "manual polygon added")
    check(lyr.commitChanges(), "manual polygon saved")
    pump(0.5)
    feats = c.store.features("buildings", "survey_id", r4["survey_id"])
    check(len(feats) == 1, "one building feature per record after manual draw")
    check(val(feats[0]["status"]) == C.STATUS_APPROVED_MANUAL, "manual drawing -> APPROVED_MANUAL")
    check(c.store.corrected(r4["survey_id"]) is not None, "manual record has corrected coordinate")
    draws = [a for a in c.store.actions(r4["survey_id"]) if a["action"] == C.ACTION_DRAW_MANUAL]
    check(draws and (bool(draws[-1]["geometry_before"]) == had_proposal), "manual draw audited with the replaced polygon")

    print("\n== pick other candidate (worker re-processes)")
    k5 = next(k for k in queue if c.records[k]["state"] == C.REVIEW and len(c.details[k]["candidates"]) > 1)
    c.select(k5)
    other = c.details[k5]["candidates"][1]["candidate_id"]
    rev_before = c.details[k5]["revision"]
    c.pick_candidate(other, k5)
    wait_until(lambda: c.details[k5]["revision"] != rev_before, 60, "select-candidate result")
    check(c.details[k5]["chosen_candidate_id"] == other, "worker re-processed from the picked candidate")
    check(c.details[k5]["decision_reason"] == "Selected by the operator", "decision records operator choice")
    check(any(a["action"] == C.ACTION_PICK_CANDIDATE for a in c.store.actions(c.records[k5]["survey_id"])),
          "pick-candidate audited")
    lon5, lat5 = c.records[k5]["longitude"], c.records[k5]["latitude"]
    check(c.candidate_at(lon5 + 10, lat5) is None, "candidate_at outside candidates returns None")

    print("\n== click-to-fix needs licensed imagery")
    try:
        c.click_to_fix([[lon5, lat5]], [], None, k5)
        check(False, "click-to-fix refused without a licensed crop")
    except UserError:
        check(True, "click-to-fix refused without a licensed crop")

    print("\n== export")
    out = os.path.join(c.batch["batch_dir"], "output", "export_test.xlsx")
    rep = c.export(out)
    approved = [k for k, r in c.records.items() if r["state"] in (C.APPROVED, C.APPROVED_MANUAL)]
    check(rep["exported"] == len(approved), "exported only approved records (%d)" % rep["exported"])
    check(os.path.isfile(rep["delivery"][0]) if rep.get("delivery") else False, "delivery GeoPackage written")
    src = openpyxl.load_workbook(workbook).active
    dst = openpyxl.load_workbook(out).active
    same = all(src.cell(r, col).value == dst.cell(r, col).value
               for r in range(1, src.max_row + 1) for col in range(1, 21))
    check(same, "source columns A-T untouched in the export")
    check(dst.max_row == src.max_row, "same number of rows")
    check(dst.cell(1, 29).value == "Corrected Latitude", "output columns appended after the template")
    check(src.merged_cells.ranges and str(list(dst.merged_cells.ranges)[0]) == "AA1:AA5", "merged template cells kept")
    hdr = {dst.cell(1, col).value: col for col in range(1, dst.max_column + 1)}
    rows = {dst.cell(r, 1).value: r for r in range(2, dst.max_row + 1)}
    sid1 = c.records[k1]["survey_id"]
    r1 = rows[sid1]
    check(dst.cell(r1, 22).value == "%.6f,%.6f" % (cc["corrected_lat"], cc["corrected_lon"]),
          "client column 'Correct and Correeted Lat &long' filled")
    check(dst.cell(r1, 21).value in (c.cfg["excel"]["remarks"].values()), "client remark uses the drop-down text")
    check(dst.cell(r1, 23).value == "tester01" and dst.cell(r1, 24).value == "Completed", "operator name + status")
    sid2 = c.records[k2]["survey_id"]
    check(dst.cell(rows[sid2], 22).value in (None, ""), "rejected record not exported")
    check(dst.cell(rows[sid1], hdr["Building Polygon ID"]).value.endswith(sid1), "polygon ID written")
    check(os.path.getsize(workbook) > 0 and openpyxl.load_workbook(workbook).active.cell(2, 22).value is None,
          "source workbook never modified")

    print("\n== audit trail")
    acts = c.store.actions()
    kinds = set(a["action"] for a in acts)
    for k in (C.ACTION_APPROVE, C.ACTION_REJECT, C.ACTION_EDIT, C.ACTION_DRAW_MANUAL, C.ACTION_PICK_CANDIDATE,
              C.ACTION_EXPORT):
        check(k in kinds, "audit contains %s" % k)

    print("\n== reopen after a worker restart (the demo path)")
    c.stop_worker()
    pump(1.0)
    check(not (c.worker and c.worker.is_running()), "worker stopped")
    bdir = c.batch["batch_dir"]
    c.open_batch(bdir)
    check(c.records[k1]["state"] == C.APPROVED and c.records[k2]["state"] == C.REJECTED, "review states survive reopening")
    k6 = next(k for k in c.queue("review") if len(c.store.features("candidates", "survey_id", c.records[k]["survey_id"])) > 1)
    c.select(k6)
    cands6 = sorted(c.store.features("candidates", "survey_id", c.records[k6]["survey_id"]), key=lambda f: val(f["rank"]))
    inner = cands6[1].geometry().pointOnSurface().asPoint()
    check(c.candidate_at(inner.x(), inner.y()) == val(cands6[1]["candidate_id"]),
          "pick tool finds candidates from review.gpkg after reopening")
    n_msgs = len(iface.bar.messages)
    c.pick_candidate(val(cands6[1]["candidate_id"]), k6)
    wait_until(lambda: len(iface.bar.messages) > n_msgs, 30, "pick error message")
    check(any("Click Process" in m[1] for m in iface.bar.messages[n_msgs:]),
          "pick before Process gives plain guidance, not a raw error")
    c.process(skip_reference_check=True)
    wait_until(lambda: c.worker_ready, 240, "worker restart")
    wait_until(lambda: not c.poll_timer.isActive() and len(c.details) >= 100, 300, "batch reload")
    pump(0.5)
    check(c.records[k1]["state"] == C.APPROVED and val(c.store.building(c.records[k1]["survey_id"])["status"]) == C.STATUS_APPROVED,
          "Process after reopening keeps approved work")
    check(c.records[k2]["state"] == C.REJECTED, "Process after reopening keeps rejected work")
    rev6 = c.details[k6]["revision"]
    c.pick_candidate(val(cands6[1]["candidate_id"]), k6)
    wait_until(lambda: c.details[k6]["revision"] != rev6, 60, "pick after reload")
    check(c.details[k6]["chosen_candidate_id"] == val(cands6[1]["candidate_id"]), "pick works after Process reload")

    c.stop_worker()
    pump(1.0)
    plugin.unload()
    return check.done()


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:]))
