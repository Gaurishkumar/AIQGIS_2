import datetime as dt

import openpyxl
import pytest

from bpa_common import constants as C
from bpa_common.entry import LocationError, parse_location
from bpa_common.excel_io import (ExcelSchemaError, cell_to_text, choose_remark, export_workbook, parse_coordinate,
                                 read_records, write_entry_workbook)

from conftest import SAMPLE_XLSX, needs_sample

HEAD = ["Survey ID", "latitude", "longitude", "Customer Name", "Building Name", "Building No", "Tower/Block",
        "Street Name", "Sector/Phase", "Area Name", "Landmark", "City", "District Name", "State Name", "pincode",
        "Standard Building Address", "Building Type", "Building Category", "Building Height(m)", "Floors",
        "Location coordinates change remark", "Correct and Correeted Lat &long", "Name", "Status", "Delivery Date"]


def make_wb(path, rows, head=HEAD):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(head)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def row(sid, lat, lon, name="Test Apt", pincode="421301"):
    return [sid, lat, lon, "Bank", name, None, None, "Station Road", None, "Kalyan", None, "Kalyan", "Thane",
            "Maharashtra", pincode, "addr", "Non-MTU", "Residential Building", 12, 4]


def test_cell_to_text_matches_pandas_str():
    assert cell_to_text(421301) == "421301"
    assert cell_to_text(19.22632) == "19.22632"
    assert cell_to_text(421301.0) == "421301"
    assert cell_to_text("  0421301 ") == "0421301"
    assert cell_to_text(None) == ""
    assert cell_to_text(True) == "TRUE"


def test_parse_coordinate():
    assert parse_coordinate("19.2263") == pytest.approx(19.2263)
    assert parse_coordinate("19,2263") == pytest.approx(19.2263)
    assert parse_coordinate("73.1°E") == pytest.approx(73.1)
    assert parse_coordinate("abc") is None
    assert parse_coordinate("") is None


def test_validation_rules(tmp_path, cfg):
    p = make_wb(str(tmp_path / "in.xlsx"), [
        row("A1", 19.2, 73.1), row("A2", 73.1, 19.2), row("A3", "x", 73.1), row("A4", 51.5, -0.12),
        row("A1", 19.3, 73.2), row(None, 19.2, 73.1), row("A5", 19.21, 73.11, name=None),
        row("A6", 19.2, 73.1, pincode="0421301")])
    recs, rep = read_records(p, cfg)
    by = {}
    for r in recs:
        by.setdefault(r["survey_id"], []).append(r)
    assert all("Duplicate Survey ID" in r["reasons"][0] for r in by["A1"])
    assert by["A2"][0]["reasons"] == ["Latitude/longitude appear swapped"]
    assert "not numeric" in by["A3"][0]["reasons"][0]
    assert "outside the configured country bounds" in by["A4"][0]["reasons"][0]
    assert by[""][0]["reasons"] == ["Survey ID is empty"]
    assert by["A5"][0]["state"] == C.IMPORTED and "Building name is empty" in by["A5"][0]["warnings"]
    assert by["A6"][0]["fields"]["pincode"] == "0421301"        # leading zero kept
    assert "Same coordinates" in " ".join(by["A6"][0]["warnings"])
    assert rep["invalid"] == 6 and rep["imported"] == 2


def test_missing_required_column(tmp_path, cfg):
    head = list(HEAD)
    head[1] = "gps_y"
    p = make_wb(str(tmp_path / "bad.xlsx"), [row("A1", 19.2, 73.1)], head)
    with pytest.raises(ExcelSchemaError, match="latitude"):
        read_records(p, cfg)


def test_loose_header_matching(tmp_path, cfg):
    head = [h.upper() if h else h for h in HEAD]
    head[0] = " survey  id "
    p = make_wb(str(tmp_path / "loose.xlsx"), [row("A1", 19.2, 73.1)], head)
    recs, _ = read_records(p, cfg)
    assert recs[0]["survey_id"] == "A1" and recs[0]["state"] == C.IMPORTED


def test_unseen_layout_sheet_and_aliases(tmp_path, cfg):
    """A new client file: other sheet name, shuffled columns, 'Lat'/'Lng' headers, extra column."""
    head = ["Customer Name", "SURVEY ID", "Building Name", "Lat", "Lng", "Street Name", "Extra notes",
            "Location coordinates change remark", "Correct and Correeted Lat &long", "Name", "Status", "Delivery Date"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey Data"
    ws.append(head)
    ws.append(["Bank", "000123", "Kalika Tower", 19.203615, 72.974116, "Kolbad Road", "x"])
    ws.append(["Bank", "B2", "Other", "18.5236° N", "73.8397° E", "FC Road", "y"])
    src = str(tmp_path / "new.xlsx")
    wb.save(src)
    recs, rep = read_records(src, cfg)
    assert rep["sheet"] == "Survey Data" and rep["notes"]
    assert [r["survey_id"] for r in recs] == ["000123", "B2"]
    assert recs[1]["latitude"] == pytest.approx(18.5236) and recs[1]["state"] == C.IMPORTED
    res = {"000123": {"status": C.STATUS_APPROVED, "corrected_lat": 19.2036, "corrected_lon": 72.9741,
                      "reviewer": "op", "shift_m": 5.0, "dist_to_polygon_m": 0.0}}
    out = str(tmp_path / "new_out.xlsx")
    rep2 = export_workbook(src, out, res, cfg)
    ws2 = openpyxl.load_workbook(out)["Survey Data"]
    assert rep2["exported"] == 1
    assert ws2.cell(2, 9).value == "19.203600,72.974100"          # client column found by header, not position
    assert ws2.cell(2, 7).value == "x" and ws2.cell(1, 13).value == "Corrected Latitude"


def test_no_usable_sheet(tmp_path, cfg):
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    wb.active.append(["foo", "bar"])
    p = str(tmp_path / "bad.xlsx")
    wb.save(p)
    with pytest.raises(ExcelSchemaError, match="no sheet with Survey ID"):
        read_records(p, cfg)


def test_remark_rules(cfg):
    rm = cfg["excel"]["remarks"]
    assert choose_remark(cfg, 0.4, 0.0) == rm["unchanged"]
    assert choose_remark(cfg, 8.0, 0.0) == rm["adjusted_near"]
    assert choose_remark(cfg, 60.0, 45.0) == rm["modified"]


def test_export_copy_only_approved(tmp_path, cfg):
    src = make_wb(str(tmp_path / "src.xlsx"), [row("A1", 19.2, 73.1), row("A2", 19.25, 73.15), row("A3", 19.3, 73.2)])
    before = open(src, "rb").read()
    results = {
        "A1": {"status": C.STATUS_APPROVED, "corrected_lat": 19.2000041, "corrected_lon": 73.1000059,
               "polygon_id": "b-A1", "area_m2": 250.26, "method": "footprint", "confidence": 0.61234,
               "route": "REVIEW", "reviewer": "op1", "reviewed_at": "2026-09-27T10:00:00", "strategy": "centroid_if_inside",
               "shift_m": 5.0, "dist_to_polygon_m": 0.0},
        "A2": {"status": C.STATUS_REJECTED, "corrected_lat": 1, "corrected_lon": 2},
        "A3": {"status": C.STATUS_PROPOSED, "corrected_lat": 1, "corrected_lon": 2},
        "ZZ": {"status": C.STATUS_APPROVED, "corrected_lat": 1, "corrected_lon": 2},
    }
    out = str(tmp_path / "out.xlsx")
    rep = export_workbook(src, out, results, cfg, exported_at=dt.datetime(2026, 9, 27, 12, 0))
    assert open(src, "rb").read() == before
    assert rep["exported"] == 1 and rep["unknown_ids"] == ["ZZ"]
    ws = openpyxl.load_workbook(out).active
    assert ws.cell(2, 22).value == "19.200004,73.100006"
    assert ws.cell(2, 21).value == cfg["excel"]["remarks"]["adjusted_near"]
    assert ws.cell(2, 23).value == "op1" and ws.cell(2, 24).value == "Completed"
    assert ws.cell(3, 22).value is None and ws.cell(4, 22).value is None
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert hdr[:25] == HEAD and hdr[25] == "Corrected Latitude"
    assert ws.cell(2, 26).value == pytest.approx(19.200004)


def test_export_refuses_to_overwrite_source(tmp_path, cfg):
    src = make_wb(str(tmp_path / "src.xlsx"), [row("A1", 19.2, 73.1)])
    with pytest.raises(ExcelSchemaError):
        export_workbook(src, src, {}, cfg)


@needs_sample
def test_client_sample_imports(cfg):
    import pandas as pd
    recs, rep = read_records(SAMPLE_XLSX, cfg)
    assert rep["total"] == 100 and rep["invalid"] == 0
    df = pd.read_excel(SAMPLE_XLSX, sheet_name="Sheet1", engine="openpyxl", dtype=str).fillna("")
    for i, r in enumerate(recs):
        for f, h in cfg["excel"]["columns"].items():
            assert r["fields"][f] == str(df.iloc[i][h]).strip()
    assert all("Don" not in k for k in recs[0]["source"])      # template control columns ignored


INDIA = {"min_lat": 6.0, "max_lat": 37.6, "min_lon": 68.0, "max_lon": 97.5}


@pytest.mark.parametrize("text,lat,lon", [
    ("19.07604, 72.87773", 19.07604, 72.87773),
    ("19.07604 72.87773", 19.07604, 72.87773),
    ("19.0760° N, 72.8777° E", 19.076, 72.8777),
    ("19°04'33.7\"N 72°52'39.8\"E", 19.0760278, 72.8777222),
    ("https://www.google.com/maps/place/X/@19.2263,73.1396,17z/data=!4m6!3m5!8m2!3d19.226321!4d73.139621",
     19.226321, 73.139621),                                            # the pin, not the map centre
    ("https://www.google.com/maps?q=19.226321,73.139621", 19.226321, 73.139621),
    ("https://www.google.com/maps/search/?api=1&query=19.2263%2C73.1396", 19.2263, 73.1396),
    ("72.87773, 19.07604", 19.07604, 72.87773),                        # swapped
])
def test_parse_location(text, lat, lon):
    r = parse_location(text, INDIA)
    assert r["lat"] == pytest.approx(lat, abs=1e-6) and r["lon"] == pytest.approx(lon, abs=1e-6)


@pytest.mark.parametrize("text", ["", "hello", "51.5, -0.12", "https://maps.app.goo.gl/abcd"])
def test_parse_location_rejects(text):
    with pytest.raises(LocationError):
        parse_location(text, INDIA)


def test_single_entries_make_a_normal_batch_workbook(tmp_path, cfg):
    entries = [{"survey_id": "SE-1", "latitude": 19.2263219, "longitude": 73.1396211, "building_name": "Kalika Tower",
                "street": "Pune Link Road", "city": "Kalyan", "pincode": "421301", "building_category": "Residential Building",
                "floors": "7", "customer_name": "Icici Bank Limited"},
               {"survey_id": "SE-2", "latitude": 19.25, "longitude": 73.15}]
    path = write_entry_workbook(cfg, entries, str(tmp_path / "single.xlsx"))
    recs, rep = read_records(path, cfg)
    assert rep["total"] == 2 and rep["invalid"] == 0 and not rep["missing_optional_columns"]
    a, b = recs
    assert a["latitude"] == pytest.approx(19.2263219) and a["fields"]["building_name"] == "Kalika Tower"
    assert a["fields"]["customer_name"] == "Icici Bank Limited" and a["state"] == C.IMPORTED
    assert "Building name is empty" in b["warnings"]
    results = {"SE-1": {"status": C.STATUS_APPROVED, "corrected_lat": 19.22633, "corrected_lon": 73.13963,
                        "reviewer": "op1", "shift_m": 2.0, "dist_to_polygon_m": 0.0}}
    out = str(tmp_path / "out.xlsx")
    assert export_workbook(path, out, results, cfg)["exported"] == 1
    ws = openpyxl.load_workbook(out).active
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col = hdr.index(cfg["excel"]["client_output_columns"]["corrected_latlon"]) + 1
    assert ws.cell(2, col).value == "19.226330,73.139630"
